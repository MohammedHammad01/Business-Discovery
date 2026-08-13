"""DOCX adapter (python-docx), imported lazily like the PDF adapter."""

from __future__ import annotations

import io

from app.schemas.canonical import SourceType
from app.services.ingestion.base import (
    ExtractorUnavailableError,
    IngestionError,
    NormalizedSource,
    extension,
)
from app.services.normalization.classifier import normalize_plain_text

_DOCX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def docx_available() -> bool:
    try:
        import docx  # noqa: F401
    except Exception:
        return False
    return True


class DocxAdapter:
    name = "docx"

    def supports(self, filename: str, content_type: str) -> bool:
        return extension(filename) == "docx" or content_type == _DOCX_CONTENT_TYPE

    def extract(self, filename: str, content_type: str, data: bytes) -> NormalizedSource:
        try:
            import docx
        except Exception as exc:  # pragma: no cover - environment dependent
            raise ExtractorUnavailableError(
                "DOCX extraction needs python-docx. Install it with `pip install python-docx`, "
                "or upload the document as TXT/PDF."
            ) from exc

        try:
            document = docx.Document(io.BytesIO(data))
        except Exception as exc:
            raise IngestionError(f"Could not read '{filename}' as a DOCX file: {exc}") from exc

        blocks = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    blocks.append(" | ".join(cells))

        content = normalize_plain_text("\n".join(blocks))
        if not content.strip():
            raise IngestionError(f"'{filename}' contains no readable text.")

        return NormalizedSource(
            source_type=SourceType.docx,
            name=filename,
            content=content,
            metadata={
                "extraction": "python-docx",
                "paragraph_count": len(document.paragraphs),
                "content_type": content_type,
            },
        )
